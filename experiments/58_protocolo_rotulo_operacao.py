#!/usr/bin/env python3
"""EXP 58 — ETAPA 1 DA REANÁLISE: o rótulo dos episódios pela OPERAÇÃO, às cegas, com controle negativo.

PRÉ-REGISTRO (20/10/2026), COMMITADO ANTES DE ENVIAR QUALQUER PERGUNTA. Gera a lista para a operação e, quando as respostas chegarem, aplica a regra fixada aqui.

POR QUÊ. Hoje "FP" quer dizer "não terminou em trip em 48 h" — não "não havia problema". Se a operação viu algo e interveio, o alarme foi útil e está sendo contado como erro.
Há pistas concretas nos dados (diário de alarmes e grade): um alarme que termina numa parada de 77 h antes do trip de 04/11/2025; pressão baixa de óleo com a máquina operando 7 h
antes de outro; HI/HIHI do termopar TC382_03_A durante dois. Mas pergunta dirigida aos FP induz resposta: daí o desenho às cegas e o controle negativo.

DESENHO (fixado antes):
  JANELAS: os 18 episódios do detector publicado desde 01/02/2025 (ledger do controle) + 9 JANELAS DE CONTROLE sorteadas (semente 20261024) em operação estável, com duração
  sorteada entre as dos 18 episódios, sem alarme do detector a menos de 48 h, a mais de 7 d de qualquer trip, >= 80% do tempo em operação estável, sem se sobrepor entre si.
  As 27 vão em ordem de data com códigos J01..J27, horário LOCAL (UTC-3) e UTC. NADA diz quais são acerto, falso positivo ou controle; a chave fica só com a equipe.
  PERGUNTA, IGUAL PARA TODAS: no intervalo [início, fim + 72 h], houve (a) ordem de serviço, (b) intervenção ou inspeção com achado, (c) parada não planejada atribuída a problema do
  equipamento, ou (d) registro de turno de anomalia — em qual subsistema e com qual evidência (nº da OS, nota de turno)?
  CATEGORIAS: U = evidência documentada de anomalia/intervenção no turbocompressor; P = parada planejada, utilidade/gás, manutenção programada SEM achado; N = nada anormal; D = sem
  registro / não sabe. Só conta U com evidência citada.
  PERGUNTAS ESPECÍFICAS, distribuídas entre os três tipos de janela (para não sinalizar quais interessam): o alarme de pressão baixa de óleo em operação em 31/03/2025 04:45 e
  02/04/2025 02:12 (locais); o motivo da parada de 31/10/2025 06:42 a 03/11/2025 12:04 (locais, 77 h); HI/HIHI do TC382_03_A em 20/10/2025 13:18 e 08/04/2026 10:21-18:53 (locais);
  o motivo das paradas que encerram as janelas de trip (o trip em si já é conhecido); e, para as janelas de controle, a mesma pergunta aberta.
EMENDA (21/10/2026, ANTES DO ENVIO; nenhuma resposta recebida): perguntas específicas 8 a 11 acrescentadas — o motivo das paradas que encerram dois FP com ΔP baixo na linha
  de balanceamento do gás de selagem (PDI-6240302, a tag do trip de 29/04/2025): 20/10/2025 13:28 e 04/04/2026 05:56 (locais); o que foi feito na manutenção de 31/10 10:24 a
  01/11 06:38 (local), entre o FP de 26/10/2025 e o trip de óleo de 04/11; e, para manter o desenho às cegas, as janelas de controle J06 e J22. A regra de decisão não muda.
EMENDA 2 (24/10/2026, ANTES DO ENVIO; nenhuma resposta recebida), vinda da anatomia dos trips a 30 s (21/10) e da triagem dos FP (22/10): complemento na pergunta 4 (o trip
  de 27/02/2025 foi vazamento confirmado ou suspeita de leitura do PDIT-6240305?) e perguntas 12 a 14: (12) o que foi feito no mancal radial LNA entre os trips de 17/03, 07/04,
  11/04 e 29/04/2025, e se a máquina foi religada sem intervenção em algum deles; (13) se os alarmes de muito alto de 27/02/2025 14:33 (PDIT-6240305) e 18/03/2025 08:16
  (TI-6240305), locais (17:33 e 11:16 UTC no diário), levaram a parada — a completeza do registro de trips; (14) se houve problema no ar de instrumentação (PI-5134001, U-5134)
  nos períodos J18 e J20 — J20 é o FP puxado pela PI_5134001 na triagem e J18 é janela de controle, para manter o desenho às cegas (dito só aqui, não no documento enviado).
  A regra de decisão não muda; as perguntas específicas são descritivas. Subcomando `pacote`: a planilha de resposta (.xlsx, com listas de escolha) em data_local/exp58_envio/.
REGRA DE DECISÃO (fixada antes):
  FP LÍQUIDO = FP - (nº de FP com U). Relatado SEMPRE junto do FP bruto; o alvo (trip) não muda.
  CONTROLE NEGATIVO: a reclassificação só é CRÍVEL se (U entre os 9 FP) - (U entre as 9 de controle) >= 3; relata-se o p de Fisher unilateral. Se a operação marcar U nas janelas de
  controle tanto quanto nos FP, o "alarme útil" não é distinguível de rotina e o FP bruto continua sendo o número.
  META EM FP LÍQUIDO = FP líquido <= 6 com reclassificação CRÍVEL (a decisão de adotar a métrica líquida é do usuário).
  SECUNDÁRIO: o subsistema citado nas U contra a família atribuída pelo EXP 55 (A1), como validação externa da camada "onde olhar".
EXPECTATIVAS: a operação responde >= 20 das 27 janelas (~50%); >= 2 FP com U (~45%); reclassificação CRÍVEL (~25%); META EM FP LÍQUIDO (~15%); a parada de 77 h antes do 04/11
tem causa ligada ao óleo ou ao turbocompressor (~35%).
LIMITES: memória e registro da operação variam com o tempo (fev/2025 é mais difícil que abr/2026); 9 contra 9 tem pouco poder; a pergunta específica pode ainda induzir; nada muda
o detector — muda a contabilidade e dá rótulo para o futuro.

Uso:  PYTHONPATH=src .venv/bin/python experiments/58_protocolo_rotulo_operacao.py gera       (lista às cegas + chave, em data_local/)
      PYTHONPATH=src .venv/bin/python experiments/58_protocolo_rotulo_operacao.py pacote     (planilha de resposta para a operação em data_local/exp58_envio/; SEM a chave)
      PYTHONPATH=src .venv/bin/python experiments/58_protocolo_rotulo_operacao.py analisa    (lê data_local/exp58_respostas.xlsx, aba "periodos", ou exp58_respostas.csv)
"""
from __future__ import annotations
import sys, warnings
from pathlib import Path
import numpy as np
import pandas as pd
from scipy import stats

warnings.filterwarnings("ignore")
HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parent / "src"))
from nv import io as nvio
from nv.causal import carrega_registro
from nv.avalia import carrega_controle
from nv.avalia.metricas import episodios_de_serie, JANELA

ROOT = HERE.parent / "data_local"
INICIO8 = pd.Timestamp("2025-02-01", tz="UTC"); FIM = pd.Timestamp("2026-04-20", tz="UTC")
SEMENTE = 20261024; N_CONTROLE = 9; LOCAL = pd.Timedelta(hours=-3)


def gera():
    g = nvio.carrega_grade(); est = nvio.mascaras(g)["estavel"]
    al_c, vig_c = carrega_controle("atual"); trips = carrega_registro()
    eps = [(a, b) for a, b in episodios_de_serie(al_c.loc[INICIO8:vig_c.index[-1]])]
    cls = ["TP" if any(a <= t and b >= t - JANELA for t in trips) else "FP" for a, b in eps]
    assert len(eps) == 18 and cls.count("FP") == 9, (len(eps), cls.count("FP"))
    dur = np.array([(b - a).total_seconds() / 3600 for a, b in eps])
    rng = np.random.default_rng(SEMENTE); cand = est[est].loc[INICIO8:FIM].index; ctrl = []
    proibido = [(a - JANELA, b + JANELA) for a, b in eps]
    while len(ctrl) < N_CONTROLE:
        s = cand[rng.integers(len(cand))]; d = float(rng.choice(dur)); e = s + pd.Timedelta(hours=d)
        if e > FIM or any(s <= y and e >= x for x, y in proibido) or any(abs((t - s).total_seconds()) < 7 * 86400 or abs((t - e).total_seconds()) < 7 * 86400 for t in trips):
            continue
        if any(s <= y and e >= x for x, y in ctrl) or float(est.loc[s:e].mean()) < 0.8:
            continue
        ctrl.append((s, e))
    linhas = [dict(inicio=a, fim=b, tipo=c) for (a, b), c in zip(eps, cls)] + [dict(inicio=a, fim=b, tipo="CONTROLE") for a, b in ctrl]
    d = pd.DataFrame(linhas).sort_values("inicio").reset_index(drop=True)
    d["codigo"] = [f"J{i + 1:02d}" for i in range(len(d))]
    d["horas"] = ((d.fim - d.inicio).dt.total_seconds() / 3600).round(1)
    for c in ("inicio", "fim"):
        d[f"{c}_local"] = (d[c] + LOCAL).dt.strftime("%d/%m/%Y %H:%M"); d[f"{c}_utc"] = d[c].dt.strftime("%d/%m/%Y %H:%M")
    d[["codigo", "inicio_local", "fim_local", "horas", "inicio_utc", "fim_utc"]].to_csv(ROOT / "exp58_janelas_para_operacao.csv", index=False)
    d[["codigo", "tipo", "inicio_utc", "fim_utc", "horas"]].to_csv(ROOT / "exp58_chave_NAO_ENVIAR.csv", index=False)
    print(f"{len(d)} janelas: {d.tipo.value_counts().to_dict()} | lista às cegas e chave gravadas em {ROOT}")
    print(d[["codigo", "inicio_local", "fim_local", "horas"]].to_string(index=False))


SUBSISTEMAS = ["mancais", "selagem", "óleo lubrificante", "vibração", "exaustão/combustão", "gás combustível", "ar de instrumentação", "outro"]
PROTOCOLO = HERE.parent / "PROTOCOLO_CONSULTA_OPERACAO_TC33003A.md"


def pacote():
    """Planilha de resposta para a operação: aba "periodos" (J01-J27 com listas de escolha) e aba "perguntas" (as específicas, lidas do protocolo). Não lê a chave."""
    import re
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    jan = pd.read_csv(ROOT / "exp58_janelas_para_operacao.csv")
    texto = PROTOCOLO.read_text(encoding="utf-8").split("## Perguntas específicas", 1)[1]
    pergs = [(int(n), re.sub(r"\*\*", "", q).strip()) for n, q in re.findall(r"^(\d+)\. (.+)$", texto, flags=re.M)]
    wb = Workbook(); ws = wb.active; ws.title = "periodos"; cab = Font(bold=True); fundo = PatternFill("solid", fgColor="F0F4F8")
    ws.append(["código", "início (local, UTC−3)", "fim (local, UTC−3)", "duração (h)", "categoria (U/P/N/D)", "subsistema", "evidência (nº da OS, nota de turno)", "comentário"])
    for r in jan.itertuples():
        ws.append([r.codigo, r.inicio_local, r.fim_local, float(r.horas), None, None, None, None])
    dv_c = DataValidation(type="list", formula1='"U,P,N,D"', allow_blank=True); dv_s = DataValidation(type="list", formula1='"' + ",".join(SUBSISTEMAS) + '"', allow_blank=True)
    ws.add_data_validation(dv_c); ws.add_data_validation(dv_s); dv_c.add(f"E2:E{len(jan) + 1}"); dv_s.add(f"F2:F{len(jan) + 1}")
    for col, w in zip("ABCDEFGH", (8, 20, 20, 11, 18, 22, 40, 40)):
        ws.column_dimensions[col].width = w
    for c in ws[1]:
        c.font = cab; c.fill = fundo; c.alignment = Alignment(wrap_text=True, vertical="top")
    wq = wb.create_sheet("perguntas"); wq.append(["nº", "pergunta", "resposta", "evidência (nº da OS, nota de turno)"])
    for n, q in pergs:
        wq.append([n, q, None, None])
    for col, w in zip("ABCD", (5, 90, 50, 35)):
        wq.column_dimensions[col].width = w
    for row in wq.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in wq[1]:
        c.font = cab; c.fill = fundo
    wi = wb.create_sheet("instrucoes")
    for linha in ("Para cada período (J01 a J27), considere o intervalo do início até 72 h depois do fim.",
                  "U: anomalia ou intervenção no turbocompressor COM registro (OS, inspeção com achado, parada não planejada por problema do equipamento, anotação de turno). Evidência obrigatória.",
                  "P: parada planejada, problema de utilidade ou gás combustível, ou manutenção programada SEM achado.", "N: nada de anormal registrado.", "D: sem registro / não é possível saber.",
                  "Horários em hora local (UTC−3). Não há resposta certa: 'nada de anormal' é tão útil quanto as outras."):
        wi.append([linha])
    wi.column_dimensions["A"].width = 140
    saida = ROOT / "exp58_envio"; saida.mkdir(exist_ok=True); arq = saida / "EXP58_respostas_TC33003A.xlsx"; wb.save(arq)
    print(f"planilha: {arq} ({len(jan)} períodos, {len(pergs)} perguntas)")


def analisa():
    ch = pd.read_csv(ROOT / "exp58_chave_NAO_ENVIAR.csv")
    xl = ROOT / "exp58_respostas.xlsx"
    if xl.exists():
        rs = pd.read_excel(xl, sheet_name="periodos").rename(columns=lambda c: {"código": "codigo", "categoria (U/P/N/D)": "categoria", "evidência (nº da OS, nota de turno)": "evidencia",
                                                                                 "comentário": "comentario"}.get(c, c))
    else:
        rs = pd.read_csv(ROOT / "exp58_respostas.csv")
    d = ch.merge(rs, on="codigo", how="left"); d["U"] = (d.categoria.astype(str).str.upper() == "U") & d.evidencia.notna() & (d.evidencia.astype(str).str.strip() != "")
    resp = int(d.categoria.notna().sum()); ufp = int(d[d.tipo == "FP"].U.sum()); uct = int(d[d.tipo == "CONTROLE"].U.sum()); utp = int(d[d.tipo == "TP"].U.sum())
    p = float(stats.fisher_exact([[ufp, 9 - ufp], [uct, 9 - uct]], alternative="greater")[1]); cred = (ufp - uct) >= 3
    fp_liq = 9 - ufp if cred else 9
    print(f"respondidas {resp}/27 | U entre FP {ufp}/9, entre controles {uct}/9, entre TP {utp}/9 | Fisher unilateral p {p:.3f} -> reclassificação {'CRÍVEL' if cred else 'NÃO crível'}")
    print(f"FP bruto 9 | FP líquido {fp_liq} -> {'META EM FP LÍQUIDO' if (cred and fp_liq <= 6) else 'meta não atingida em FP líquido'}")
    print(d[["codigo", "tipo", "categoria", "subsistema", "evidencia"]].to_string(index=False))


if __name__ == "__main__":
    {"gera": gera, "pacote": pacote, "analisa": analisa}[sys.argv[1] if len(sys.argv) > 1 else "gera"]()
